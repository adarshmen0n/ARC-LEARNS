from pypdf import PdfReader
from docx import Document
import csv
import os


def extract_pdf_text(file_path: str):

    reader = PdfReader(file_path)

    text = ""

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


def extract_docx_text(file_path: str):

    document = Document(file_path)

    text = ""

    for paragraph in document.paragraphs:

        if paragraph.text.strip():
            text += paragraph.text + "\n"

    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:
                row_text.append(cell.text)

            text += " | ".join(row_text) + "\n"

    return text


def extract_txt_text(file_path: str):

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as file:

        return file.read()


def extract_csv_text(file_path: str):

    text = ""

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore",
        newline=""
    ) as file:

        reader = csv.reader(file)

        for row in reader:

            text += " | ".join(row) + "\n"

    return text


def extract_text(file_path: str):

    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":
        return extract_pdf_text(file_path)

    elif extension == ".docx":
        return extract_docx_text(file_path)

    elif extension == ".txt":
        return extract_txt_text(file_path)

    elif extension == ".csv":
        return extract_csv_text(file_path)

    else:
        raise ValueError(
            "Unsupported file type. "
            "Supported formats: PDF, DOCX, TXT, CSV."
        )